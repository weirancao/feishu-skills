"""技能：根据会议记录要点生成结构化会议纪要（Word 文档）"""
import os
import time

from .base import Skill
import config


class MeetingMinutesSkill(Skill):
    name = "meeting_minutes"
    description = (
        "根据会议主题、参与人和讨论要点，整理成一份结构化的会议纪要文档(.docx)并发送到当前会话。"
        "当用户要求写会议纪要、整理会议记录时调用。"
    )
    parameters = {
        "topic": {"type": "string", "description": "会议主题"},
        "participants": {"type": "array", "items": {"type": "string"}, "description": "参会人"},
        "notes": {"type": "array", "items": {"type": "string"}, "description": "会议讨论要点/结论/待办，每项一句话"},
    }
    required = ["topic", "notes"]

    def run(self, args: dict, ctx):
        from openai import OpenAI
        from docx import Document

        topic = args.get("topic", "")
        participants = args.get("participants", [])
        notes = args.get("notes", [])
        cfg = config.MODELS[config.DEFAULT_MODEL]
        client = OpenAI(api_key=cfg["api_key"], base_url=cfg["base_url"])
        resp = client.chat.completions.create(
            model=cfg["model"],
            messages=[
                {"role": "system", "content": "你是行政助理，把会议要点整理成结构化纪要：分为「会议结论」「待办事项」「风险与跟进」三部分，每条用短句，专业简洁。"},
                {"role": "user", "content": f"会议主题：{topic}\n参会人：{'、'.join(participants)}\n讨论要点：\n" + "\n".join(f"- {n}" for n in notes)},
            ],
        )
        text = resp.choices[0].message.content.strip()

        doc = Document()
        doc.add_heading(f"会议纪要：{topic}", 0)
        doc.add_paragraph(f"参会人：{'、'.join(participants) if participants else '（未提供）'}")
        doc.add_heading("内容", 1)
        for line in text.splitlines():
            if line.strip():
                doc.add_paragraph(line)

        out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "output")
        os.makedirs(out_dir, exist_ok=True)
        file_name = f"MEETING_{time.strftime('%Y%m%d_%H%M%S')}.docx"
        path = os.path.join(out_dir, file_name)
        doc.save(path)

        ctx.send_file(path, file_name)
        return "会议纪要已生成并发送到会话，请用户查收。"
