"""技能：生成可编辑 Word 文档并发送到当前会话"""
import os
import time

from .base import Skill


class CreateDocxSkill(Skill):
    name = "create_docx"
    description = (
        "根据内容大纲生成一份可编辑的Word文档(.docx)并自动发送到当前会话。"
        "当用户要求写报告、写文档、生成Word文件时调用。"
    )
    parameters = {
        "title": {"type": "string", "description": "文档标题"},
        "sections": {
            "type": "array",
            "description": "章节列表，元素为 {heading: 小节标题, paragraphs: [段落文本]}",
            "items": {
                "type": "object",
                "properties": {
                    "heading": {"type": "string", "description": "小节标题"},
                    "paragraphs": {"type": "array", "items": {"type": "string"}, "description": "段落文本"},
                },
                "required": ["heading", "paragraphs"],
            },
        },
    }
    required = ["title", "sections"]

    def run(self, args: dict, ctx):
        from docx import Document

        doc = Document()
        doc.add_heading(args.get("title", "文档"), 0)
        for sec in args.get("sections", []):
            doc.add_heading(sec.get("heading", ""), 1)
            for p in sec.get("paragraphs", []):
                doc.add_paragraph(p)

        out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "output")
        os.makedirs(out_dir, exist_ok=True)
        file_name = f"DOC_{time.strftime('%Y%m%d_%H%M%S')}.docx"
        path = os.path.join(out_dir, file_name)
        doc.save(path)

        ctx.send_file(path, file_name)
        return "Word 文档已生成并发送到会话，请用户查收。"
